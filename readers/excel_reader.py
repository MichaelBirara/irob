from pathlib import Path
import openpyxl
import xlrd

class ExcelReader:
    def __init__(self, file_path):
        self.file_path = Path(file_path)

    def read(self):
        if not self.file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {self.file_path}")
        
        ext = self.file_path.suffix.lower()
        ir_data = {}

        if ext == '.xlsx':
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = []
                headers = []
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i == 0:
                        headers = [str(h) if h is not None else f"col_{j}" for j, h in enumerate(row)]
                        continue
                    if not any(row):
                        continue
                    row_dict = {headers[j]: val for j, val in enumerate(row) if j < len(headers)}
                    rows.append(row_dict)
                ir_data[sheet_name] = rows
        elif ext == '.xls':
            wb = xlrd.open_workbook(self.file_path)
            for sheet_name in wb.sheet_names():
                sheet = wb.sheet_by_name(sheet_name)
                rows = []
                headers = []
                for i in range(sheet.nrows):
                    row_values = sheet.row_values(i)
                    if i == 0:
                        headers = [str(h) if h is not None else f"col_{j}" for j, h in enumerate(row_values)]
                        continue
                    if not any(row_values):
                        continue
                    row_dict = {headers[j]: val for j, val in enumerate(row_values) if j < len(headers)}
                    rows.append(row_dict)
                ir_data[sheet_name] = rows
        else:
            raise ValueError(f"Unsupported Excel format: {ext}")

        return ir_data
